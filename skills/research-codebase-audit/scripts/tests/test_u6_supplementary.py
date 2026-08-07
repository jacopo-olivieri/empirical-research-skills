"""U6b supplementary wave, late observations, bC, and production drills."""

import json
import subprocess
import sys

import openpyxl
import pytest

import regbuild as rb


certify = rb.load_script("certify_stage")
export = rb.load_script("export_xlsx")
lint = rb.load_script("lint_registers")
assembler = rb.load_script("assemble_boundary")

pytestmark = pytest.mark.u6

CERTIFY = rb.SCRIPTS_DIR / "certify_stage.py"


def cli(root, command, *args):
    return subprocess.run(
        [sys.executable, str(CERTIFY), command, "--package-root", str(root),
         *map(str, args)], capture_output=True, text=True)


def footer(rows=()):
    return "\n### Footer dispositions\n\n" + rb.md_table(lint.FOOTER_COLS, list(rows))


def code_shard(ledger_rows, footer_rows=()):
    return (
        rb.register_text("Recheck ledger", rb.CODE_LEDGER_COLS, list(ledger_rows))
        + "\n### Witness outcomes\n\n"
        + rb.md_table(rb.WITNESS_OUTCOME_COLS, [])
        + "\n### Verification records\n\nNo verification records.\n"
        + footer(footer_rows)
    )


def claims_shard(ledger_rows, footer_rows=()):
    return (
        rb.register_text("Recheck ledger", rb.LEDGER_COLS, list(ledger_rows))
        + "\nCoverage: every assigned claim was dispositioned.\n"
        + footer(footer_rows)
    )


def zero_supplementary_plan():
    return (rb.recheck_plan_text("code", [], [])
            + f"\n{lint.SUPPLEMENTARY_EMPTY}\n")


def lo_artifact(rows=(), dispositions=()):
    text = "# Late observations — code\n\n"
    text += rb.md_table(lint.LO_COLS, list(rows)) if rows else lint.LO_EMPTY + "\n"
    text += "\n## Dispositions\n\n"
    normalized = [
        row if len(row) == 3 else [row[0], "pending", row[1]]
        for row in dispositions
    ]
    text += (rb.md_table(lint.LO_DISPOSITION_COLS, normalized)
             if dispositions else lint.LO_DISPOSITIONS_EMPTY + "\n")
    return text


def make_wave(tmp_path, *, discovery=True, include_supplementary_ledger=True):
    root = tmp_path / "package"
    root.mkdir()
    (root / "source.py").write_text("VALUE = 1\n", encoding="utf-8")
    a = rb.AuditDir(root)
    a.write_manifest(mode="code_errors_only", scope_exclusions=[], off_limits=[])
    a.write_register("code_error_register.md", rb.ERROR_COLS, [
        rb.error_row("E-0100", status="candidate", severity="2")])
    initialized = cli(root, "init")
    assert initialized.returncode == 0, initialized.stdout + initialized.stderr

    a.write("plans/code_error_review_plan.md", rb._code_b1_plan())
    a.write("plans/code_error_second_read_plan.md", "# Code second-read plan\n")
    a.write("_run/detector_mapping.md", rb.detector_mapping_artifact([]))
    main_inventory = [("E-0100", "issue-flagged", "source.py")]
    main_clusters = [("K1", "main", "E-0100", "`audit/_code_error_recheck/k1.md`")]
    a.write("plans/code_error_recheck_plan.md",
            rb.recheck_plan_text("code", main_inventory, main_clusters))
    main_ledger = rb.code_ledger_row(
        "E-0100", verdict="confirmed_error", proposed_status="confirmed",
        proposed_severity="2", witness_ids="—")
    main_footer = (["OBS-0001", "candidate", "", "fresh issue", ""]
                   if discovery else None)
    a.write("_code_error_recheck/k1.md",
            code_shard([main_ledger], [main_footer] if main_footer else []))
    before = [rb.error_row("E-0100", status="candidate", severity="2")]
    after_b6a = [rb.error_row("E-0100", status="confirmed", severity="2")]
    if discovery:
        after_b6a.append(rb.error_row("E-8000", status="candidate", severity="2"))
    a.write_register("_run/snapshots/code_b6a/code_error_register.md",
                     rb.ERROR_COLS, before)
    a.write_register("_run/snapshots/code_b6b/code_error_register.md",
                     rb.ERROR_COLS, after_b6a)
    summary = (
        "# Code recheck summary\n\nSplits declared: 0\nMerges declared: 0\n"
        f"Discoveries declared: C=0; O=0; E={1 if discovery else 0}\n"
    )
    if discovery:
        summary += "\naudit/_code_error_recheck/k1.md#OBS-0001 | candidate:E-8000\n"
    a.write("code_error_recheck_summary.md", summary)
    if discovery:
        supp_inventory = [("E-8000", "b6a-discovery", "source.py")]
        supp_clusters = [("KS1", "supplementary", "E-8000",
                          "`audit/_code_error_recheck_supplementary/k1.md`")]
        supp_plan = (rb.recheck_plan_text("code", supp_inventory, supp_clusters)
                     + "\nDeclared supplementary discovery range: E-8000–E-8000\n")
    else:
        supp_plan = zero_supplementary_plan()
    a.write("plans/code_error_supplementary_recheck_plan.md", supp_plan)
    a.write("_run/code_b6a/dismissal_receipts.md",
            "# Dismissal receipts\n\nNo mapped not_error dismissal receipts were required.\n")
    a.write("_run/code_b6a/witness_outcomes.md", assembler.render([], []))

    final = list(after_b6a)
    shard = None
    if discovery:
        supp_ledger = rb.code_ledger_row(
            "E-8000", verdict="confirmed_error", proposed_status="confirmed",
            proposed_severity="2", witness_ids="—")
        shard = a.write("_code_error_recheck_supplementary/k1.md",
                        code_shard([supp_ledger] if include_supplementary_ledger else []))
        final[-1] = rb.error_row("E-8000", status="confirmed", severity="2")
    a.write_register("code_error_register.md", rb.ERROR_COLS, final)
    a.write("code_error_supplementary_recheck_summary.md",
            "# Supplementary recheck summary\n")
    a.write("late_observations_code.md", lo_artifact())
    a.write("_run/code_b6b/dismissal_receipts.md",
            "# Supplementary dismissal receipts\n\n"
            "No supplementary dismissal receipts were required.\n")
    a.write("_run/code_b6b/witness_outcomes.md",
            "# Supplementary witness outcomes\n\n"
            "No supplementary mapped witness outcomes.\n")
    return root, a, shard


def make_claims_zero_wave(tmp_path):
    root = tmp_path / "claims-package"
    root.mkdir()
    (root / "source.py").write_text("VALUE = 1\n", encoding="utf-8")
    a = rb.AuditDir(root)
    a.write_manifest(
        mode="replication", scope_exclusions=[], off_limits=[],
        allocation_override={"purpose": "development", "allocation": []})
    before_claim = rb.claims_row(
        "C-0100", status="candidate", severity="2", issue="possible mismatch",
        outputs="O-0100")
    final_claim = rb.claims_row(
        "C-0100", status="inconsistent", severity="2", issue="confirmed mismatch",
        outputs="O-0100")
    output = rb.output_row("O-0100", claims="C-0100", status="mapped")
    a.write_register("claims_register.md", rb.CLAIMS_COLS, [before_claim])
    a.write_register("output_register.md", rb.OUTPUT_COLS, [output])
    initialized = cli(root, "init")
    assert initialized.returncode == 0, initialized.stdout + initialized.stderr
    a.write_claims_plan()
    a.write("plans/claims_second_read_plan.md", "# Claims second-read plan\n")
    inventory = [("C-0100", "issue-flagged", "paper.tex")]
    clusters = [("K1", "main", "C-0100", "`audit/_recheck/k1.md`")]
    a.write("plans/claims_recheck_plan.md",
            rb.recheck_plan_text("claims", inventory, clusters))
    a.write("_recheck/k1.md", claims_shard([
        rb.ledger_row(
            "C-0100", verdict="substantiated", change="set status=inconsistent",
            severity="2", note="confirmed mismatch")]))
    a.write_register("_run/snapshots/claims_b6a/claims_register.md",
                     rb.CLAIMS_COLS, [before_claim])
    a.write_register("_run/snapshots/claims_b6a/output_register.md",
                     rb.OUTPUT_COLS, [output])
    a.write_register("_run/snapshots/claims_b6b/claims_register.md",
                     rb.CLAIMS_COLS, [final_claim])
    a.write_register("_run/snapshots/claims_b6b/output_register.md",
                     rb.OUTPUT_COLS, [output])
    a.write_register("claims_register.md", rb.CLAIMS_COLS, [final_claim])
    a.write_register("output_register.md", rb.OUTPUT_COLS, [output])
    a.write("claims_recheck_summary.md",
            "# Claims recheck summary\n\nSplits declared: 0\nMerges declared: 0\n"
            "Discoveries declared: C=0; O=0; E=0\n")
    a.write("plans/claims_supplementary_recheck_plan.md",
            rb.recheck_plan_text("claims", [], [])
            + f"\n{lint.SUPPLEMENTARY_EMPTY}\n")
    a.write("claims_supplementary_recheck_summary.md",
            "# Supplementary recheck summary\n")
    a.write("late_observations_claims.md",
            "# Late observations — claims\n\nNo late observations.\n\n"
            "## Dispositions\n\nNo dispositions.\n")
    return root, a


def certify_wave(root, a, *, discovery=True):
    for stage in ("code_b5", "code_b6a", "code_b5s", "code_b6b"):
        started = cli(root, "start", "--stage", stage)
        assert started.returncode == 0, started.stdout + started.stderr
        if stage == "code_b5":
            recorded = cli(root, "set-shard", "--stage", stage, "--shard",
                           "audit/_code_error_recheck/k1.md", "--status", "done")
            assert recorded.returncode == 0, recorded.stdout + recorded.stderr
        if stage == "code_b5s" and discovery:
            recorded = cli(root, "set-shard", "--stage", stage, "--shard",
                           "audit/_code_error_recheck_supplementary/k1.md",
                           "--status", "done")
            assert recorded.returncode == 0, recorded.stdout + recorded.stderr
        finished = cli(root, "finish", "--stage", stage, "--outcome", "done")
        assert finished.returncode == 0, finished.stdout + finished.stderr


def certify_to_b5s(root, *, discovery=True):
    for stage in ("code_b5", "code_b6a", "code_b5s"):
        started = cli(root, "start", "--stage", stage)
        assert started.returncode == 0, started.stdout + started.stderr
        if stage == "code_b5":
            recorded = cli(root, "set-shard", "--stage", stage, "--shard",
                           "audit/_code_error_recheck/k1.md", "--status", "done")
            assert recorded.returncode == 0, recorded.stdout + recorded.stderr
        if stage == "code_b5s" and discovery:
            recorded = cli(
                root, "set-shard", "--stage", stage, "--shard",
                "audit/_code_error_recheck_supplementary/k1.md", "--status", "done")
            assert recorded.returncode == 0, recorded.stdout + recorded.stderr
        finished = cli(root, "finish", "--stage", stage, "--outcome", "done")
        assert finished.returncode == 0, finished.stdout + finished.stderr


def test_stage_keys_replace_b6_and_include_single_cycle_and_bc():
    for stages in (certify.FULL_STAGES, certify.CODE_ONLY_STAGES):
        assert "code_b6" not in stages
        assert stages.index("code_b6a") < stages.index("code_b5s") < stages.index("code_b6b")
        assert stages.count("bC") == 1
    assert "claims_b6" not in certify.FULL_STAGES
    assert certify.FULL_STAGES.index("claims_b6a") < certify.FULL_STAGES.index("claims_b5s") < certify.FULL_STAGES.index("claims_b6b")


def test_b5s_zero_work_certifies_from_plan_without_dummy_shard(tmp_path):
    root, a, _ = make_wave(tmp_path, discovery=False)
    certify_wave(root, a, discovery=False)
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    assert manifest["stages"]["code_b5s"]["status"] == "done"
    assert manifest["stages"]["code_b5s"]["shards"] == {}
    verified = cli(root, "verify-run")
    assert verified.returncode == 0, verified.stdout + verified.stderr


def test_nonempty_supplementary_wave_certifies_and_verify_run_passes(tmp_path):
    root, a, _ = make_wave(tmp_path)
    certify_wave(root, a)
    verified = cli(root, "verify-run")
    assert verified.returncode == 0, verified.stdout + verified.stderr


def test_claims_zero_work_supplementary_wave_certifies_and_verifies(tmp_path):
    root, a = make_claims_zero_wave(tmp_path)
    for stage in ("claims_b5", "claims_b6a", "claims_b5s", "claims_b6b"):
        assert cli(root, "start", "--stage", stage).returncode == 0
        if stage == "claims_b5":
            recorded = cli(root, "set-shard", "--stage", stage, "--shard",
                           "audit/_recheck/k1.md", "--status", "done")
            assert recorded.returncode == 0, recorded.stderr
        finished = cli(root, "finish", "--stage", stage, "--outcome", "done")
        assert finished.returncode == 0, finished.stdout + finished.stderr
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    assert manifest["stages"]["claims_b5s"]["shards"] == {}
    verified = cli(root, "verify-run")
    assert verified.returncode == 0, verified.stdout + verified.stderr


def test_tier1_supplementary_closure_cli_refuses_deleted_ledger_and_verify(tmp_path):
    root, a, shard = make_wave(tmp_path)
    certify_to_b5s(root)
    shard.write_text(code_shard([]), encoding="utf-8")
    assert cli(root, "start", "--stage", "code_b6b").returncode == 0
    refused = cli(root, "finish", "--stage", "code_b6b", "--outcome", "done")
    assert refused.returncode != 0
    assert "has 0 ledger rows" in refused.stderr
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    manifest["stages"]["code_b6b"]["status"] = "done"
    certify.write_manifest_atomic(root, manifest)
    verified = cli(root, "verify-run")
    assert verified.returncode != 0 and "code_b6b" in verified.stderr


def test_tier1_supplementary_closure_check_has_teeth(tmp_path, monkeypatch):
    root, a, shard = make_wave(tmp_path)
    certify_to_b5s(root)
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    original_ledgers = lint._all_recheck_ledgers(lint.Lint(), a.audit, "code", True)
    shard.write_text(code_shard([]), encoding="utf-8")
    live = lint.Lint()
    lint.stage_b6b(live, a.audit, "code", manifest)
    assert any("has 0 ledger rows" in error for error in live.errors)
    monkeypatch.setattr(
        lint, "_all_recheck_ledgers",
        lambda _lint, _audit, _stream, supplementary:
        original_ledgers if supplementary else [],
    )
    broken = lint.Lint()
    lint.stage_b6b(broken, a.audit, "code", manifest)
    assert not any("has 0 ledger rows" in error for error in broken.errors)


def test_tier1_mixed_elimination_lint_and_verify_refuse(tmp_path):
    root, a, shard = make_wave(tmp_path)
    certify_wave(root, a)
    shard.write_text(shard.read_text().replace(
        "guard changes the selected sample", "MIXED"), encoding="utf-8")
    direct = rb.lint(a, "b6b-code")
    assert direct.returncode != 0 and "MIXED cannot survive" in direct.stdout
    verified = cli(root, "verify-run")
    assert verified.returncode != 0 and "code_b6b" in verified.stderr


def test_tier1_mixed_elimination_check_has_teeth(tmp_path, monkeypatch):
    root, a, shard = make_wave(tmp_path)
    certify_to_b5s(root)
    shard.write_text(shard.read_text().replace(
        "guard changes the selected sample", "MIXED"), encoding="utf-8")
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    live = lint.Lint()
    lint.stage_b6b(live, a.audit, "code", manifest)
    assert any("MIXED cannot survive" in error for error in live.errors)
    monkeypatch.setattr(lint, "_reject_mixed_cells", lambda *_args: None)
    broken = lint.Lint()
    lint.stage_b6b(broken, a.audit, "code", manifest)
    assert not any("MIXED cannot survive" in error for error in broken.errors)


def make_bc(tmp_path):
    root = tmp_path / "bc-package"
    root.mkdir()
    (root / "source.py").write_text("VALUE = 1\n", encoding="utf-8")
    a = rb.AuditDir(root)
    a.write_manifest(mode="code_errors_only", scope_exclusions=[], off_limits=[])
    before = [rb.error_row("E-0100", status="confirmed", severity="2")]
    added = rb.error_row("E-8001", status="confirmed", severity="2")
    a.write_register("code_error_register.md", rb.ERROR_COLS, before)
    assert cli(root, "init").returncode == 0
    a.write("plans/code_error_review_plan.md", rb._code_b1_plan())
    a.write("plans/code_error_second_read_plan.md", "# Code second-read plan\n")
    a.write("plans/code_error_supplementary_recheck_plan.md", zero_supplementary_plan())
    a.write("_run/detector_mapping.md", rb.detector_mapping_artifact([]))
    a.write_register("_run/snapshots/bC/code_error_register.md", rb.ERROR_COLS, before)
    a.write_register("code_error_register.md", rb.ERROR_COLS, before + [added])
    a.write("late_observations_code.md", lo_artifact(
        [["LO-E-0001", "audit/_code_error_recheck_supplementary/k1.md#OBS-0001",
          "source.py:1", "late defect"]],
        [["LO-E-0001", "minted:BC-0001"]]))
    a.write("_run/snapshots/bC/late_observations_code.md", lo_artifact(
        [["LO-E-0001", "audit/_code_error_recheck_supplementary/k1.md#OBS-0001",
          "source.py:1", "late defect"]],
        [["LO-E-0001", "pending"]]))
    payload = json.dumps(dict(zip(rb.ERROR_COLS, added)), sort_keys=True,
                         separators=(",", ":"))
    plan = (
        "# Late-observation corrections\n\nDeclared bC range: E-8001–E-8001\n\n"
        + rb.md_table(lint.BC_PLAN_COLS, [[
            "BC-0001", "LO-E-0001", "code_error", "new_row", "E-8001",
            payload, "—"]]))
    a.write("plans/late_observation_corrections.md", plan)
    return root, a, before, added


def test_tier1_bc_delta_cli_and_verify_refuse_undeclared_cell(tmp_path):
    root, a, before, added = make_bc(tmp_path)
    tampered = rb.error_row("E-0100", status="confirmed", severity="2",
                            desc="undeclared edit")
    a.write_register("code_error_register.md", rb.ERROR_COLS, [tampered, added])
    assert cli(root, "start", "--stage", "bC").returncode == 0
    refused = cli(root, "finish", "--stage", "bC", "--outcome", "done")
    assert refused.returncode != 0 and "undeclared bC cell change" in refused.stderr
    a.write_register("code_error_register.md", rb.ERROR_COLS, before + [added])
    finished = cli(root, "finish", "--stage", "bC", "--outcome", "done")
    assert finished.returncode == 0, finished.stderr
    a.write_register("code_error_register.md", rb.ERROR_COLS, [tampered, added])
    verified = cli(root, "verify-run")
    assert verified.returncode != 0 and "bC" in verified.stderr


def test_tier1_bc_delta_check_has_teeth(tmp_path, monkeypatch):
    _root, a, _before, added = make_bc(tmp_path)
    tampered = rb.error_row(
        "E-0100", status="confirmed", severity="2", desc="undeclared edit")
    a.write_register("code_error_register.md", rb.ERROR_COLS, [tampered, added])
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    live = lint.Lint()
    lint.stage_bC(live, a.audit, manifest)
    assert any("undeclared bC cell change" in error for error in live.errors)
    monkeypatch.setattr(lint, "_check_bc_declared_cells", lambda *_args: None)
    broken = lint.Lint()
    lint.stage_bC(broken, a.audit, manifest)
    assert not any("undeclared bC cell change" in error for error in broken.errors)


def set_supplementary_late_observation(a, *, source=None):
    shard = a.audit / "_code_error_recheck_supplementary/k1.md"
    text = shard.read_text(encoding="utf-8")
    text = text.replace(
        rb.md_table(lint.FOOTER_COLS, []),
        rb.md_table(lint.FOOTER_COLS, [[
            "OBS-0001", "candidate", "", "second-order defect", ""]]))
    shard.write_text(text, encoding="utf-8")
    source = source or "audit/_code_error_recheck_supplementary/k1.md#OBS-0001"
    a.write("code_error_supplementary_recheck_summary.md",
            "# Supplementary recheck summary\n\n"
            "audit/_code_error_recheck_supplementary/k1.md#OBS-0001 | "
            "late_observation:LO-E-0001\n")
    a.write("late_observations_code.md", lo_artifact(
        [["LO-E-0001", source, "source.py:1", "second-order defect"]],
        [["LO-E-0001", "pending"]]))


def test_late_observation_footer_reconciliation_accepts_exact_join(tmp_path):
    root, a, _ = make_wave(tmp_path)
    set_supplementary_late_observation(a)
    certify_wave(root, a)


def test_late_observation_footer_reconciliation_refuses_wrong_join(tmp_path):
    root, a, _ = make_wave(tmp_path)
    set_supplementary_late_observation(
        a, source="audit/_code_error_recheck_supplementary/other.md#OBS-0001")
    certify_to_b5s(root)
    assert cli(root, "start", "--stage", "code_b6b").returncode == 0
    refused = cli(root, "finish", "--stage", "code_b6b", "--outcome", "done")
    assert refused.returncode != 0
    assert "does not join exactly" in refused.stderr


def test_output_discovery_lifecycle_accepts_each_branch_and_refuses_overlap(tmp_path):
    summary = tmp_path / "summary.md"
    orphan = rb.output_row("O-8000", claims="", status="orphan")
    structural = lint.Lint()
    lint._output_discovery_lifecycle(
        structural, summary, "", [orphan], {"O-8000"})
    assert not structural.errors

    mapped = rb.output_row("O-8000", claims="C-8000", status="inconsistent")
    table = rb.md_table(lint.OUTPUT_DISCOVERY_COLS, [[
        "O-8000", "C-8000", "substantiated", "inconsistent"]])
    joined = lint.Lint()
    lint._output_discovery_lifecycle(
        joined, summary, table, [mapped], {"O-8000"})
    assert not joined.errors

    overlap = lint.Lint()
    lint._output_discovery_lifecycle(
        overlap, summary, table, [orphan], {"O-8000"})
    assert any("exactly one" in error for error in overlap.errors)


def test_detector_split_descendant_materiality_requires_structured_ruling(tmp_path):
    _root, a, shard = make_wave(tmp_path)
    a.write("_run/detector_mapping.md", rb.detector_mapping_artifact([
        ("DU-aaaaaaaaaaaa", "E-0100", "new_candidate")]))
    lineage = rb.md_table(lint.LINEAGE_COLS, [[
        "E-0100", "E-8000", "DU", "DU-aaaaaaaaaaaa", "DUW-000000000001"]])
    a.write("code_error_recheck_summary.md",
            "# Code recheck summary\n\nSplits declared: 1\nMerges declared: 0\n"
            "Discoveries declared: C=0; O=0; E=0\n\n" + lineage)
    blocked = rb.code_ledger_row(
        "E-8000", verdict="blocked", proposed_status="blocked",
        proposed_severity="2", note="needs another source")
    shard.write_text(code_shard([blocked]), encoding="utf-8")
    bad = lint.Lint()
    lint._check_detector_materiality(bad, a.audit)
    assert any("materiality_reassessment" in error for error in bad.errors)
    good_row = rb.code_ledger_row(
        "E-8000", verdict="blocked", proposed_status="blocked",
        proposed_severity="2",
        note="[materiality_reassessment] severity=2; basis=downstream sample impact")
    shard.write_text(code_shard([good_row]), encoding="utf-8")
    good = lint.Lint()
    lint._check_detector_materiality(good, a.audit)
    assert not good.errors

    a.write_register("code_error_register.md", rb.ERROR_COLS, [
        rb.error_row("E-0100", status="confirmed", severity="2"),
        rb.error_row("E-8000", status="confirmation_needed", severity="2")])
    confirmation = rb.code_ledger_row(
        "E-8000", verdict="confirmation_needed", proposed_status="confirmation_needed",
        proposed_severity="3",
        note="[materiality_reassessment] severity=2; basis=uncertainty caps the applied ruling")
    shard.write_text(code_shard([confirmation]), encoding="utf-8")
    applied = lint.Lint()
    lint._check_detector_materiality(applied, a.audit)
    assert not applied.errors
    shard.write_text(code_shard([rb.code_ledger_row(
        "E-8000", verdict="confirmation_needed", proposed_status="confirmation_needed",
        proposed_severity="3",
        note="[materiality_reassessment] severity=3; basis=left provisional")]),
        encoding="utf-8")
    unapplied = lint.Lint()
    lint._check_detector_materiality(unapplied, a.audit)
    assert any("was not applied" in error for error in unapplied.errors)


def test_bc_rejects_illegal_recorded_disposition_transition(tmp_path):
    _root, a, _before, _added = make_bc(tmp_path)
    snapshot = a.audit / "_run/snapshots/bC/late_observations_code.md"
    snapshot.write_text(snapshot.read_text().replace(
        "pending", "qa_commissioned:QA-0001"), encoding="utf-8")
    result = rb.lint(a, "bC")
    assert result.returncode != 0
    assert "illegal disposition transition" in result.stdout


def test_bc_new_row_comparison_accepts_only_authorized_b8_rewrite():
    headers = [
        "Error ID", "Error Description", "Error Description Original", "Status"]
    payload = {
        "Error ID": "E-8001", "Error Description": "worker wording",
        "Error Description Original": "", "Status": "confirmed",
    }
    rewritten = {
        "Error ID": "E-8001", "Error Description": "author wording",
        "Error Description Original": "worker wording", "Status": "confirmed",
    }
    pairs = [("Error Description", "Error Description Original")]
    assert lint._bc_new_row_matches(payload, rewritten, headers, pairs)
    rewritten["Status"] = "blocked"
    assert not lint._bc_new_row_matches(payload, rewritten, headers, pairs)


def test_b6b_reverification_uses_frozen_pre_bc_register(tmp_path):
    root, a, _ = make_wave(tmp_path, discovery=False)
    certify_wave(root, a, discovery=False)
    frozen_rows = [rb.error_row("E-0100", status="confirmed", severity="2")]
    frozen_cols, frozen_rewritten = rb.rewrite_pass_cols(
        rb.ERROR_COLS, frozen_rows, ["Error Description", "Why It Matters"])
    a.write_register("_run/snapshots/bC/code_error_register.md",
                     frozen_cols, frozen_rewritten)
    live_cols, live_rewritten = rb.rewrite_pass_cols(
        rb.ERROR_COLS, frozen_rows + [
            rb.error_row("E-9000", status="confirmed", severity="2")],
        ["Error Description", "Why It Matters"])
    a.write_register("code_error_register.md", live_cols, live_rewritten)
    result = rb.lint(a, "b6b-code")
    assert result.returncode == 0, result.stdout + result.stderr


def test_duplicate_chain_must_resolve_to_present_terminal_row(tmp_path):
    path = tmp_path / "register.md"
    good_rows = [
        rb.error_row("E-0100", status="confirmed", severity="2"),
        rb.error_row("E-8000", status="duplicate_of:E-0100", severity=""),
    ]
    good = lint.Lint()
    lint._check_duplicate_chains(
        good, path, good_rows, rb.ERROR_COLS, "Error ID")
    assert not good.errors
    bad_rows = good_rows + [
        rb.error_row("E-8001", status="duplicate_of:E-8000", severity="")]
    bad = lint.Lint()
    lint._check_duplicate_chains(
        bad, path, bad_rows, rb.ERROR_COLS, "Error ID")
    assert any("not resolved" in error for error in bad.errors)


def test_supplementary_boundary_projects_split_witnesses_to_b6b_homes(tmp_path):
    a = rb.AuditDir(tmp_path)
    a.write_manifest(stages={
        "code_b6a": {"status": "done", "retries": 0, "shards": {}},
        "code_b5s": {"status": "running", "retries": 0, "shards": {}},
    })
    mappings = [
        ("DU-aaa111", "E-0100", "new_candidate"),
        ("DU-bbb222", "E-0100", "new_candidate"),
    ]
    a.write("_run/detector_mapping.md", rb.detector_mapping_artifact(mappings))
    lineage = rb.md_table(lint.LINEAGE_COLS, [
        ["E-0100", "E-8000", "DU", "DU-aaa111", "DUW-000000000001"],
        ["E-0100", "E-8001", "DU", "DU-bbb222", "DUW-000000000002"],
    ])
    a.write("code_error_recheck_summary.md", "# Main merge\n\n" + lineage)
    inventory = [
        ("E-8000", "split-descendant", "source.py"),
        ("E-8001", "split-descendant", "source.py"),
    ]
    clusters = [("KS1", "split", "E-8000; E-8001",
                 "`audit/_code_error_recheck_supplementary/k1.md`")]
    a.write("plans/code_error_supplementary_recheck_plan.md",
            rb.recheck_plan_text("code", inventory, clusters))
    ledgers = [
        rb.code_ledger_row("E-8000", witness_ids="DUW-000000000001"),
        rb.code_ledger_row("E-8001", witness_ids="DUW-000000000002"),
    ]
    outcomes = [
        rb.witness_outcome_row("DU", "DU-aaa111", "DUW-000000000001"),
        rb.witness_outcome_row("DU", "DU-bbb222", "DUW-000000000002"),
    ]
    shard = a.write(
        "_code_error_recheck_supplementary/k1.md",
        rb.register_text("Recheck ledger", rb.CODE_LEDGER_COLS, ledgers)
        + "\n### Witness outcomes\n\n"
        + rb.md_table(rb.WITNESS_OUTCOME_COLS, outcomes)
        + "\n### Verification records\n\nNo verification records.\n"
        + footer())
    assert rb.lint(a, "b5s-code", shard=shard).returncode == 0
    rows = [
        rb.error_row("E-0100", status="not_error", severity=""),
        rb.error_row("E-8000", status="confirmed", severity="2"),
        rb.error_row("E-8001", status="confirmed", severity="2"),
    ]
    a.write_register("code_error_register.md", rb.ERROR_COLS, rows)
    a.write("_run/code_b6b/dismissal_receipts.md",
            "# Supplementary dismissal receipts\n\n"
            "No supplementary dismissal receipts were required.\n")
    post, dismissals, dispositions = assembler.assemble(
        a.audit, a.audit / "code_error_register.md", supplementary=True)
    assert len(post) == 2 and not dismissals
    assert set(dispositions) == {"E-8000", "E-8001"}
    a.write("_run/code_b6b/witness_outcomes.md",
            assembler.render(post, dismissals, supplementary=True))
    assert assembler.check_boundary(a.audit, supplementary=True).name == "witness_outcomes.md"


@pytest.mark.parametrize("before,after,allowed", [
    ("pending", "acknowledged_unverified", True),
    ("pending", "qa_closed:QA-0001:conclusive", False),
    ("qa_commissioned:QA-0001", "qa_closed:QA-0001:inconclusive", True),
    ("qa_commissioned:QA-0001", "qa_closed:QA-0002:conclusive", False),
    ("qa_closed:QA-0001:conclusive", "minted:BC-0001", True),
    ("minted:BC-0001", "acknowledged_unverified", False),
])
def test_late_observation_transition_matrix(before, after, allowed):
    assert lint.valid_lo_transition(before, after) is allowed


def test_b9_export_omits_late_observation_sheets_but_keeps_coverage_md(tmp_path):
    """U17 inversion: the LO sheets are absent from the workbook while the
    coverage md still carries the explicit-absence values."""
    a = rb.make_b9(tmp_path, error_rows=[rb.error_row("E-0100")],
                   mode="code_errors_only")
    workbook = openpyxl.load_workbook(a.audit / "code_review.xlsx", read_only=True)
    assert "Late observations (unverified)" not in workbook.sheetnames
    assert "Late observation coverage" not in workbook.sheetnames
    coverage = (a.audit / "_run/late_observation_coverage.md").read_text()
    assert "| not recorded | none recorded |" in coverage
    assert rb.lint(a, "b9").returncode == 0


def test_b9_records_blocked_collection_as_degraded_not_zero(tmp_path):
    a = rb.make_b9(tmp_path, error_rows=[rb.error_row("E-0100")],
                   mode="code_errors_only")
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    manifest["stages"] = {
        "code_b6b": {"status": "blocked", "retries": 0, "shards": {}}}
    (a.audit / "_run/manifest.json").write_text(
        json.dumps(manifest), encoding="utf-8")
    (a.audit / "late_observations_code.md").unlink()
    regenerated = rb.run_script(
        "export_xlsx.py", "--audit-dir", a.audit,
        "--mode", "code_errors_only", "-o", a.audit / "code_review.xlsx")
    assert regenerated.returncode == 0, regenerated.stdout + regenerated.stderr
    result = rb.lint(a, "b9")
    assert result.returncode == 0, result.stdout + result.stderr
    coverage = (a.audit / "_run/late_observation_coverage.md").read_text()
    assert "| code | yes | blocked | degraded |" in coverage


def test_b6a_cli_refuses_manifest_without_b6a_stage_key(tmp_path):
    a = rb.AuditDir(tmp_path)
    a.write_manifest(stages={"claims_b5": {"status": "done", "retries": 0,
                                           "shards": {}}})
    result = rb.lint(a, "b6a-claims")
    assert result.returncode != 0
    assert "must restart" in result.stdout


def test_close_run_is_the_completion_report_gate_for_pending_dispositions(tmp_path):
    root, a, _ = make_wave(tmp_path)
    set_supplementary_late_observation(a)
    certify_wave(root, a)
    refused = cli(root, "close-run")
    assert refused.returncode != 0
    assert "pending" in refused.stderr and "LO-E-0001" in refused.stderr
    a.write("late_observations_code.md", lo_artifact(
        [["LO-E-0001", "audit/_code_error_recheck_supplementary/k1.md#OBS-0001",
          "source.py:1", "second-order defect"]],
        [["LO-E-0001", "pending", "acknowledged_unverified"]]))
    closed = cli(root, "close-run")
    assert closed.returncode == 0, closed.stdout + closed.stderr


def test_b9_accepts_pending_dispositions_without_refusing(tmp_path):
    """b9 accepts pending dispositions: the pending row's surviving home is the
    `late_observations_code.md` collection artifact, which the export leaves
    untouched (U17 removed its workbook sheet, not the artifact)."""
    a = rb.make_b9(tmp_path, error_rows=[rb.error_row("E-0100")],
                   mode="code_errors_only")
    a.write("late_observations_code.md", lo_artifact(
        [["LO-E-0001", "audit/_code_error_recheck_supplementary/k1.md#OBS-0001",
          "source.py:1", "late defect"]],
        [["LO-E-0001", "pending"]]))
    regenerated = rb.run_script(
        "export_xlsx.py", "--audit-dir", a.audit,
        "--mode", "code_errors_only", "-o", a.audit / "code_review.xlsx")
    assert regenerated.returncode == 0, regenerated.stdout + regenerated.stderr
    result = rb.lint(a, "b9")
    assert result.returncode == 0, result.stdout + result.stderr
    artifact = (a.audit / "late_observations_code.md").read_text()
    assert "LO-E-0001" in artifact and "| pending |" in artifact


def test_b6b_refuses_dangling_duplicate_targets_in_claims_and_output(tmp_path):
    root, a = make_claims_zero_wave(tmp_path)
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    dangling_claim = rb.claims_row(
        "C-0100", status="duplicate_of:C-9999", severity="", issue="",
        outputs="O-0100")
    a.write_register("claims_register.md", rb.CLAIMS_COLS, [dangling_claim])
    a.write_register(
        "_run/snapshots/claims_b6b/claims_register.md", rb.CLAIMS_COLS,
        [dangling_claim])
    claims_lint = lint.Lint()
    lint.stage_b6b(claims_lint, a.audit, "claims", manifest)
    assert any("absent/self target C-9999" in error for error in claims_lint.errors)

    (tmp_path / "output-leg").mkdir()
    root, a = make_claims_zero_wave(tmp_path / "output-leg")
    manifest = json.loads((a.audit / "_run/manifest.json").read_text())
    dangling_output = rb.output_row(
        "O-0100", claims="C-0100", status="duplicate_of:O-9999")
    a.write_register("output_register.md", rb.OUTPUT_COLS, [dangling_output])
    a.write_register(
        "_run/snapshots/claims_b6b/output_register.md", rb.OUTPUT_COLS,
        [dangling_output])
    output_lint = lint.Lint()
    lint.stage_b6b(output_lint, a.audit, "claims", manifest)
    assert any("absent/self target O-9999" in error for error in output_lint.errors)


def test_tier1_mixed_register_cell_refused_by_b6b_lint_and_verify_run(tmp_path):
    root, a, _ = make_wave(tmp_path)
    certify_wave(root, a)
    register = a.audit / "code_error_register.md"
    lines = register.read_text(encoding="utf-8").splitlines()
    for index, line in enumerate(lines):
        if line.startswith("| E-8000"):
            cells = line.split("|")
            cells[3] = " MIXED "
            lines[index] = "|".join(cells)
    register.write_text("\n".join(lines) + "\n", encoding="utf-8")
    direct = rb.lint(a, "b6b-code")
    assert direct.returncode != 0 and "MIXED cannot survive" in direct.stdout
    verified = cli(root, "verify-run")
    assert verified.returncode != 0 and "code_b6b" in verified.stderr


def test_b3d_check_is_pinned_to_frozen_b3b_snapshot(tmp_path):
    import test_u3_adjudication as u3

    root, a = u3._completed_b3d_b6_run(tmp_path)
    register = a.audit / "code_error_register.md"
    frozen = a.audit / "_run/snapshots/code_b3b/code_error_register.md"
    frozen.parent.mkdir(parents=True, exist_ok=True)
    frozen.write_text(register.read_text(encoding="utf-8"), encoding="utf-8")
    baseline = rb.run_script(
        "build_detector_mapping.py", root, "--audit-dir", a.audit, "--check")
    assert baseline.returncode == 0, baseline.stdout + baseline.stderr

    later = rb.error_row("E-7500", status="confirmed", severity="2")
    register.write_text(
        register.read_text(encoding="utf-8").rstrip()
        + "\n| " + " | ".join(later) + " |\n", encoding="utf-8")
    pinned = rb.run_script(
        "build_detector_mapping.py", root, "--audit-dir", a.audit, "--check")
    assert pinned.returncode == 0, pinned.stdout + pinned.stderr

    # Split descendants are outside the b3d stage era: a descendant row in the
    # frozen image is an undeclared row even when split lineage names it.
    descendant = rb.error_row("E-7001", status="confirmed", severity="2")
    frozen_text = frozen.read_text(encoding="utf-8")
    frozen.write_text(
        frozen_text.rstrip() + "\n| " + " | ".join(descendant) + " |\n",
        encoding="utf-8")
    summary = a.audit / "code_error_recheck_summary.md"
    summary.write_text(
        summary.read_text(encoding="utf-8") + "\n### Split lineage\n\n"
        + rb.md_table(lint.LINEAGE_COLS, [[
            "E-7000", "E-7001", "DU", "DU-aaaaaaaaaaaa", "DUW-000000000001"]]),
        encoding="utf-8")
    excluded = rb.run_script(
        "build_detector_mapping.py", root, "--audit-dir", a.audit, "--check")
    assert excluded.returncode != 0 and "E-7001" in excluded.stderr

    # Before the snapshot exists the check falls back to live canon and the
    # later-minted row still refuses (the pre-snapshot window is unchanged).
    frozen.unlink()
    fallback = rb.run_script(
        "build_detector_mapping.py", root, "--audit-dir", a.audit, "--check")
    assert fallback.returncode != 0 and "E-7500" in fallback.stderr


def test_documented_contract_names_full_ledger_and_mech_fields():
    registers = (rb.SKILL_DIR / "references/registers.md").read_text()
    worker = (rb.SKILL_DIR / "references/prompts/recheck-cluster-worker.md").read_text()
    assert "code_error_supplementary_recheck_plan.md" in registers
    assert "path/to/shard.md#OBS-####" in registers
    for field in ("Mech Class", "Mech Object", "Mech Relation", "Mech Expected", "Mech Actual"):
        assert field in registers and field in worker
    full = " | ".join(rb.CODE_LEDGER_COLS)
    assert full in registers and full in worker
